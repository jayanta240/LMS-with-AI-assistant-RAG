import io
import re
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from services.auth_dependency import get_current_user
from services.auth_service import hash_password
from services.department_db import get_departments
from services.postgres_db import get_db_connection
from services.user_db import create_user, get_user_by_email
from services.company_db import create_company_registration

router = APIRouter()

EMPLOYEE_HEADERS = [
    "employee_id", "first_name", "middle_name", "last_name", "gender",
    "date_of_birth", "email", "password", "mobile_number", "employment_type",
    "employment_status", "date_of_joining", "date_of_confirmation", "date_of_exit",
    "designation", "department_name", "department_head_email",
]
DEPARTMENT_HEAD_HEADERS = ["name", "email", "password", "department_name"]
COMPANY_HEADERS = [
    "company_name", "industry_type", "company_type", "gst_number", "cin_number",
    "company_description", "contact_person_name", "contact_designation", "contact_email",
    "contact_mobile", "admin_name", "admin_email", "admin_password", "admin_mobile",
    "address_line_1", "address_line_2", "city", "state", "country", "pin_code",
    "plan_name", "license_type", "start_date", "expiry_date", "trial",
    "max_active_users", "max_courses", "storage_limit", "billing_cycle",
]
GSTIN_PATTERN = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
CIN_PATTERN = re.compile(r"^[LU][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6}$")
PLAN_NAMES = {"basic", "standard", "premium", "enterprise"}
LICENSE_TYPES = {"named user", "concurrent user", "enterprise", "perpetual", "subscription"}
BILLING_CYCLES = {"monthly", "quarterly", "half-yearly", "yearly", "one-time"}


def _normalise_header(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


async def _read_rows(file: UploadFile, allowed_roles):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an Excel .xlsx file.")
    if file.content_type not in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
        "application/vnd.ms-excel",
        "",
    }:
        # Browsers and proxies vary the MIME type; extension and workbook parsing are authoritative.
        pass
    contents = await file.read()
    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded Excel file is empty.")
    try:
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        sheet = workbook.active
        raw = sheet.iter_rows(values_only=True)
        headers_row = next(raw, None)
        if not headers_row:
            raise HTTPException(status_code=400, detail="The Excel sheet has no header row.")
        headers = [_normalise_header(v) for v in headers_row]
        if len(headers) != len(set(headers)):
            raise HTTPException(status_code=400, detail="The Excel file contains duplicate column headings.")
        missing = [h for h in allowed_roles if h not in headers]
        if missing:
            raise HTTPException(status_code=400, detail="Missing required columns: " + ", ".join(missing))
        rows = []
        for sheet_row, values in enumerate(raw, start=2):
            if not values or all(v is None or str(v).strip() == "" for v in values):
                continue
            row = {headers[i]: _cell_text(values[i]) if i < len(values) else "" for i in range(len(headers))}
            row["_row_number"] = sheet_row
            rows.append(row)
            if len(rows) > 500:
                raise HTTPException(status_code=400, detail="A maximum of 500 records can be uploaded per file.")
        workbook.close()
        if not rows:
            raise HTTPException(status_code=400, detail="The Excel file has headers but no data rows.")
        return rows
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file. Please use the provided .xlsx template.")


@router.get("/api/bulk-registration/template")
def download_bulk_template(
    role: str = Query(..., pattern="^(employee|department_head|company)$"),
    current_user=Depends(get_current_user),
):
    actor_role = current_user.get("role")
    if role == "employee" and actor_role not in {"company_admin", "department_head"}:
        raise HTTPException(status_code=403, detail="You cannot download the employee template.")
    if role == "department_head" and actor_role != "company_admin":
        raise HTTPException(status_code=403, detail="Only Company Admin can download the Department Head template.")
    if role == "company" and actor_role != "super_admin":
        raise HTTPException(status_code=403, detail="Only Super Admin can download the company template.")

    headers = EMPLOYEE_HEADERS if role == "employee" else DEPARTMENT_HEAD_HEADERS if role == "department_head" else COMPANY_HEADERS
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Companies" if role == "company" else "Employees" if role == "employee" else "Department Heads"
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[letter].width = min(max(max_len + 3, 16), 28)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{role}_bulk_upload_template.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/api/bulk-registration/employees")
async def bulk_register_employees(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
):
    role = current_user.get("role")
    if role not in {"company_admin", "department_head"}:
        raise HTTPException(status_code=403, detail="Only Company Admin or Department Head can bulk-register employees.")
    company_id = current_user.get("company_id")
    if not company_id:
        raise HTTPException(status_code=403, detail="Your account is not associated with a company.")

    rows = await _read_rows(file, EMPLOYEE_HEADERS)
    department_rows = get_departments(company_id)
    departments_by_name = {str(r[2]).strip().casefold(): int(r[0]) for r in department_rows}
    heads_by_email = {}
    with get_db_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT id, email, department_id FROM users WHERE company_id = %s AND role = 'department_head'",
                (company_id,),
            )
            for head_id, head_email, head_department_id in cursor.fetchall():
                heads_by_email[str(head_email or "").strip().casefold()] = (int(head_id), head_department_id)

    created, failed = [], []
    seen_emails, seen_ids = set(), set()
    for row in rows:
        number = row["_row_number"]
        try:
            first = row.get("first_name", "").strip()
            middle = row.get("middle_name", "").strip()
            last = row.get("last_name", "").strip()
            name = " ".join(part for part in (first, middle, last) if part)
            email = row.get("email", "").strip().lower()
            employee_id = row.get("employee_id", "").strip()
            password = row.get("password", "")
            gender = row.get("gender", "").strip().title()
            dob = row.get("date_of_birth", "").strip()
            mobile = row.get("mobile_number", "").strip()
            employment_type = row.get("employment_type", "").strip().title()
            employment_status = row.get("employment_status", "").strip().title()
            doj = row.get("date_of_joining", "").strip()
            confirmation = row.get("date_of_confirmation", "").strip()
            exit_date = row.get("date_of_exit", "").strip()
            designation = row.get("designation", "").strip()
            if not all([first, last, email, employee_id, password, gender, dob, mobile, employment_type, employment_status, doj, designation]):
                raise ValueError("Missing one or more required fields.")
            if email in seen_emails:
                raise ValueError("Duplicate email in this Excel file.")
            if employee_id.casefold() in seen_ids:
                raise ValueError("Duplicate employee ID in this Excel file.")
            seen_emails.add(email)
            seen_ids.add(employee_id.casefold())
            if len(password) < 8:
                raise ValueError("Password must contain at least 8 characters.")
            dept_name = row.get("department_name", "").strip()
            if role == "department_head":
                department_id = current_user.get("department_id")
                head_id = current_user.get("user_id")
            else:
                department_id = departments_by_name.get(dept_name.casefold())
                if not department_id:
                    raise ValueError(f"Department '{dept_name}' was not found in your company.")
                head_email = row.get("department_head_email", "").strip().casefold()
                head_info = heads_by_email.get(head_email)
                if not head_info or int(head_info[1] or 0) != department_id:
                    raise ValueError("Department Head email does not belong to the selected department.")
                head_id = head_info[0]
            # Reuse the single-record registration path so all existing validations,
            # profile fields, hashing, notifications and department scope stay consistent.
            result = await _create_employee_from_bulk_row(
                current_user=current_user,
                employee_id=employee_id,
                employee_name=name,
                first_name=first,
                middle_name=middle,
                last_name=last,
                gender=gender,
                date_of_birth=dob,
                email=email,
                password=password,
                confirm_password=password,
                mobile_number=mobile,
                employment_type=employment_type,
                employment_status=employment_status,
                date_of_joining=doj,
                date_of_confirmation=confirmation,
                date_of_exit=exit_date,
                designation=designation,
                department_id=department_id,
                department_head_id=head_id,
            )
            created.append({"row": number, "name": name, "email": email, "employee_id": result.get("employee_id", employee_id)})
        except Exception as exc:
            detail = getattr(exc, "detail", None) or str(exc) or "Registration failed."
            failed.append({"row": number, "email": row.get("email", ""), "reason": str(detail)})
    return {"success": len(failed) == 0, "created_count": len(created), "failed_count": len(failed), "created": created, "failed": failed}


async def _create_employee_from_bulk_row(**values):
    # Import at call time to avoid coupling route initialization order.
    from main import create_employee_profile
    return await create_employee_profile(
        employee_id=values["employee_id"],
        employee_name=values["employee_name"],
        first_name=values["first_name"],
        middle_name=values["middle_name"],
        last_name=values["last_name"],
        gender=values["gender"],
        date_of_birth=values["date_of_birth"],
        email=values["email"],
        password=values["password"],
        confirm_password=values["confirm_password"],
        mobile_number=values["mobile_number"],
        employment_type=values["employment_type"],
        employment_status=values["employment_status"],
        date_of_joining=values["date_of_joining"],
        date_of_confirmation=values["date_of_confirmation"],
        date_of_exit=values["date_of_exit"],
        designation=values["designation"],
        department_id=values["department_id"],
        department_head_id=values["department_head_id"],
        profile_photo=None,
        current_user=values["current_user"],
    )


@router.post("/api/bulk-registration/department-heads")
async def bulk_register_department_heads(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
):
    if current_user.get("role") != "company_admin":
        raise HTTPException(status_code=403, detail="Only Company Admin can bulk-register Department Heads.")
    company_id = current_user.get("company_id")
    if not company_id:
        raise HTTPException(status_code=403, detail="Your account is not associated with a company.")
    rows = await _read_rows(file, DEPARTMENT_HEAD_HEADERS)
    department_rows = get_departments(company_id)
    departments_by_name = {str(r[2]).strip().casefold(): int(r[0]) for r in department_rows}
    created, failed = [], []
    seen_emails = set()
    for row in rows:
        number = row["_row_number"]
        name = row.get("name", "").strip()
        email = row.get("email", "").strip().lower()
        password = row.get("password", "")
        dept_name = row.get("department_name", "").strip()
        try:
            if not all([name, email, password, dept_name]):
                raise ValueError("Name, email, password, and department_name are required.")
            if email in seen_emails:
                raise ValueError("Duplicate email in this Excel file.")
            seen_emails.add(email)
            if len(password) < 8:
                raise ValueError("Password must contain at least 8 characters.")
            department_id = departments_by_name.get(dept_name.casefold())
            if not department_id:
                raise ValueError(f"Department '{dept_name}' was not found in your company.")
            if get_user_by_email(email):
                raise ValueError("Email already exists.")
            user_id = create_user(
                name=name,
                email=email,
                password_hash=hash_password(password),
                role="department_head",
                company_id=company_id,
                department_id=department_id,
            )
            created.append({"row": number, "name": name, "email": email, "user_id": user_id})
        except Exception as exc:
            detail = getattr(exc, "detail", None) or str(exc) or "Registration failed."
            failed.append({"row": number, "email": email, "reason": str(detail)})
    return {"success": len(failed) == 0, "created_count": len(created), "failed_count": len(failed), "created": created, "failed": failed}


@router.post("/api/bulk-registration/companies")
async def bulk_register_companies(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
):
    if current_user.get("role") != "super_admin":
        raise HTTPException(status_code=403, detail="Only Super Admin can bulk-register companies.")

    rows = await _read_rows(file, COMPANY_HEADERS)
    created, failed = [], []
    seen_company_names, seen_admin_emails, seen_contact_emails = set(), set(), set()

    for row in rows:
        number = row["_row_number"]
        data = {key: row.get(key, "").strip() for key in COMPANY_HEADERS}
        try:
            required = [
                "company_name", "industry_type", "contact_person_name", "contact_designation",
                "contact_email", "contact_mobile", "admin_name", "admin_email", "admin_password",
                "address_line_1", "city", "state", "country", "pin_code", "plan_name",
                "license_type", "start_date", "expiry_date", "max_active_users",
            ]
            missing = [key for key in required if not data.get(key)]
            if missing:
                raise ValueError("Missing required fields: " + ", ".join(missing))
            company_key = data["company_name"].casefold()
            admin_email = data["admin_email"].casefold()
            contact_email = data["contact_email"].casefold()
            if company_key in seen_company_names:
                raise ValueError("Duplicate company name in this Excel file.")
            if admin_email in seen_admin_emails:
                raise ValueError("Duplicate Company Admin email in this Excel file.")
            if contact_email in seen_contact_emails:
                raise ValueError("Duplicate primary contact email in this Excel file.")
            seen_company_names.add(company_key)
            seen_admin_emails.add(admin_email)
            seen_contact_emails.add(contact_email)

            email_pattern = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
            if not email_pattern.match(admin_email) or not email_pattern.match(contact_email):
                raise ValueError("Enter valid contact and Company Admin email addresses.")
            if len(data["admin_password"]) < 8:
                raise ValueError("Company Admin password must contain at least 8 characters.")
            if data["gst_number"] and not GSTIN_PATTERN.fullmatch(data["gst_number"].upper()):
                raise ValueError("GST Number must be a valid 15-character GSTIN.")
            if data["cin_number"] and not CIN_PATTERN.fullmatch(data["cin_number"].upper()):
                raise ValueError("CIN Number must be a valid 21-character CIN.")
            if data["plan_name"].casefold() not in PLAN_NAMES:
                raise ValueError("Plan Name must be Basic, Standard, Premium, or Enterprise.")
            if data["license_type"].casefold() not in LICENSE_TYPES:
                raise ValueError("Invalid License Type.")
            if data["billing_cycle"] and data["billing_cycle"].casefold() not in BILLING_CYCLES:
                raise ValueError("Invalid Billing Cycle.")
            start_date = datetime.strptime(data["start_date"], "%Y-%m-%d").date()
            expiry_date = datetime.strptime(data["expiry_date"], "%Y-%m-%d").date()
            if expiry_date < start_date:
                raise ValueError("Expiry Date cannot be before Start Date.")
            try:
                max_active_users = int(data["max_active_users"])
                if max_active_users < 1:
                    raise ValueError()
            except ValueError:
                raise ValueError("Max Active Users must be a whole number greater than zero.")
            max_courses = None
            if data["max_courses"]:
                try:
                    max_courses = int(data["max_courses"])
                    if max_courses < 0:
                        raise ValueError()
                except ValueError:
                    raise ValueError("Max Courses must be a non-negative whole number.")
            trial_value = data["trial"].casefold() in {"true", "yes", "1", "y"}
            if data["trial"] and data["trial"].casefold() not in {"true", "yes", "1", "y", "false", "no", "0", "n"}:
                raise ValueError("Trial must be TRUE or FALSE.")

            data.update({
                "contact_email": contact_email,
                "admin_email": admin_email,
                "gst_number": data["gst_number"].upper() or None,
                "cin_number": data["cin_number"].upper() or None,
                "plan_name": data["plan_name"].title(),
                "license_type": data["license_type"],
                "start_date": start_date,
                "expiry_date": expiry_date,
                "max_active_users": max_active_users,
                "max_courses": max_courses,
                "trial": trial_value,
            })
            company_id, admin_id = create_company_registration(
                data, hash_password(data["admin_password"])
            )
            created.append({
                "row": number,
                "name": data["company_name"],
                "email": admin_email,
                "company_id": company_id,
                "admin_id": admin_id,
            })
        except Exception as exc:
            detail = getattr(exc, "detail", None) or str(exc) or "Registration failed."
            failed.append({"row": number, "email": data.get("admin_email", ""), "reason": str(detail)})

    return {
        "success": len(failed) == 0,
        "created_count": len(created),
        "failed_count": len(failed),
        "created": created,
        "failed": failed,
    }
